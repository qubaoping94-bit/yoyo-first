"""Audit editable PPTX object coverage and optional rendered-PDF text visibility."""
import argparse
import json
import re
from pathlib import Path


def normalized(value):
    return re.sub(r"\s+", "", str(value or ""))


def audit(folder, rendered_pdf=None):
    from pptx import Presentation
    import fitz

    manifest = json.loads((folder / "deck-manifest.json").read_text(encoding="utf-8"))
    slides = manifest["slides"]
    pptx = Presentation(folder / "presentation-editable.pptx")
    errors = []
    if len(pptx.slides) != len(slides):
        errors.append(f"editable PPTX pages {len(pptx.slides)} != {len(slides)}")
    pdf = fitz.open(rendered_pdf) if rendered_pdf else None
    if pdf and len(pdf) != len(slides):
        errors.append(f"editable rendered PDF pages {len(pdf)} != {len(slides)}")
    for index, spec in enumerate(slides[:len(pptx.slides)]):
        page = pptx.slides[index]
        object_text = normalized("".join(shape.text for shape in page.shapes if shape.has_text_frame))
        visible_text = normalized(pdf[index].get_text()) if pdf and index < len(pdf) else None
        fields = [spec.get(k, "") for k in ("title", "title_accent", "subtitle", "conclusion")] + spec.get("items", [])
        for field in fields:
            if not field:
                continue
            if normalized(field) not in object_text:
                errors.append(f"{index+1:02d} {spec['id']}: editable object missing {field[:20]}")
            if visible_text is not None and normalized(field) not in visible_text:
                errors.append(f"{index+1:02d} {spec['id']}: rendered PDF missing {field[:20]}")
        if normalized(spec.get("notes", "")) not in normalized(page.notes_slide.notes_text_frame.text):
            errors.append(f"{index+1:02d} {spec['id']}: notes missing")
    result = {"pages": len(slides), "rendered_pdf_checked": bool(pdf), "errors": errors, "editable_pass": not errors, "pixel_identical_to_visual_version": False, "content_status": manifest["status"]}
    (folder / "editable-audit.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    if pdf:
        pdf.close()
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return result


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("folder", type=Path)
    p.add_argument("--rendered-pdf", type=Path)
    a = p.parse_args()
    raise SystemExit(0 if audit(a.folder, a.rendered_pdf)["editable_pass"] else 1)
