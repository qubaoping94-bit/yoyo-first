"""Render a built HTML deck to evidence PNGs, PDF, image-fidelity PPTX and optional editable PPTX."""
import argparse
import base64
import hashlib
import io
import json
import os
import re
from pathlib import Path


def export(html_path, output, editable=False, allow_hold=False, browser_path=None, from_frames=False):
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    manifest_path = html_path.with_name("deck-manifest.json")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest["status"] == "CONTENT_HOLD" and not allow_hold:
        raise ValueError("CONTENT_HOLD: export requires explicit --allow-hold for preview-only output")
    slides = manifest["slides"]
    output.mkdir(parents=True, exist_ok=True)
    frames = output / "frames"
    frames.mkdir(exist_ok=True)
    screenshot_paths = []
    if from_frames:
        for i in range(len(slides)):
            matches = list(frames.glob(f"slide-{i+1:02d}-*.png")) or list(frames.glob(f"slide-{i+1:02d}.png"))
            if len(matches) != 1:
                raise ValueError(f"Expected one captured frame for slide {i+1:02d}; found {len(matches)}")
            with Image.open(matches[0]) as img:
                if img.size != (1920, 1080):
                    raise ValueError(f"Wrong frame size: {matches[0]} {img.size}")
            screenshot_paths.append(matches[0])
        capture_report = frames / "capture-report.json"
        if capture_report.exists() and json.loads(capture_report.read_text(encoding="utf-8")).get("status") != "PASS":
            raise ValueError("Captured frames have unresolved visual issues")
    else:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            kwargs = {"headless": True}
            if browser_path:
                kwargs["executable_path"] = str(browser_path)
            browser = p.chromium.launch(**kwargs)
            page = browser.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=1, reduced_motion="reduce")
            errors = []
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.goto(html_path.resolve().as_uri())
            page.wait_for_load_state("load")
            assert page.locator(".slide").count() == len(slides), "HTML/manifest page count mismatch"
            for i, slide in enumerate(slides):
                page.evaluate("i => show(i)", i)
                page.evaluate("document.fonts.ready")
                frame = frames / f"slide-{i+1:02d}.png"
                page.screenshot(path=str(frame), animations="disabled")
                screenshot_paths.append(frame)
            browser.close()
        if errors:
            raise RuntimeError("Browser page errors: " + " | ".join(errors))

    images = [Image.open(path).convert("RGB") for path in screenshot_paths]
    images[0].save(output / "presentation.pdf", "PDF", resolution=144, save_all=True, append_images=images[1:])
    for img in images:
        img.close()

    pptx = Presentation()
    pptx.slide_width, pptx.slide_height = Inches(13.333333), Inches(7.5)
    blank = pptx.slide_layouts[6]
    for spec, frame in zip(slides, screenshot_paths):
        slide = pptx.slides.add_slide(blank)
        slide.shapes.add_picture(str(frame), 0, 0, width=pptx.slide_width, height=pptx.slide_height)
        slide.notes_slide.notes_text_frame.text = spec.get("notes", "")
    pptx.save(output / "presentation.pptx")

    if editable:
        html_text = html_path.read_text(encoding="utf-8")
        edit = Presentation()
        edit.slide_width, edit.slide_height = pptx.slide_width, pptx.slide_height
        blank = edit.slide_layouts[6]
        for index, spec in enumerate(slides, 1):
            slide = edit.slides.add_slide(blank)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(244, 241, 232)
            spine = slide.shapes.add_shape(1, 0, 0, Inches(.19), edit.slide_height)
            spine.fill.solid(); spine.fill.fore_color.rgb = RGBColor(230, 59, 46)
            spine.line.fill.background()
            def box(text, x, y, w, h, size, color, bold=False):
                shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
                tf = shape.text_frame; tf.word_wrap = True
                para = tf.paragraphs[0]; para.text = text
                para.font.name = "Microsoft YaHei"; para.font.size = Pt(size); para.font.bold = bold
                para.font.color.rgb = RGBColor(*color)
            display_title = spec["title"] + ("\n" + spec["title_accent"] if spec.get("title_accent") else "")
            box(display_title, .55, .65, 6.2 if spec.get("image") else 12.1, 1.75, 32, (23, 24, 21), True)
            if spec.get("subtitle"):
                box(spec["subtitle"], .6, 2.38, 11.9, .55, 16, (104, 106, 100))
            items = spec.get("items", [])
            if items:
                columns = 1 if spec.get("image") else 2
                rows = (len(items) + columns - 1) // columns
                start = 3.05 if spec.get("subtitle") else 2.68
                end = 6.42 if spec.get("conclusion") else 6.95
                pitch = (end - start) / max(1, rows)
                max_length = max(map(len, items))
                size = 15 if max_length > 75 else 17 if max_length > 45 else 18
                for n, item in enumerate(items):
                    col = n % columns; row = n // columns
                    box(f"{n+1:02d}  {item}", .65 + col*6.05, start + row*pitch, 5.7, pitch-.08, size, (23, 24, 21), True)
            if spec.get("conclusion"):
                box(spec["conclusion"], .65, 6.55, 12.0, .5, 14, (230, 59, 46), True)
            if spec.get("image"):
                section = re.search(r'<section class="slide" data-slide-id="' + re.escape(spec["id"]) + r'".*?</section>', html_text, re.S)
                encoded = re.search(r'<img[^>]+src="data:image/[^;]+;base64,([^"]+)"', section.group(0)) if section else None
                if encoded:
                    image_buffer = io.BytesIO()
                    Image.open(io.BytesIO(base64.b64decode(encoded.group(1)))).save(image_buffer, format="PNG")
                    image_buffer.seek(0)
                    slide.shapes.add_picture(image_buffer, Inches(7.2), Inches(2.0), width=Inches(5.4))
            box(f"{index:02d} / {len(slides):02d}  ·  可编辑内容版（非像素保真）", .6, 7.18, 11.8, .23, 9, (104, 106, 100))
            slide.notes_slide.notes_text_frame.text = spec.get("notes", "")
        edit.save(output / "presentation-editable.pptx")

    report = {"status": manifest["status"], "page_count": len(slides), "outputs": {}}
    for name in ["presentation.pdf", "presentation.pptx"] + (["presentation-editable.pptx"] if editable else []):
        path = output / name
        report["outputs"][name] = {"sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "bytes": path.stat().st_size, "mode": "editable-content-not-pixel-identical" if name == "presentation-editable.pptx" else "visual-fidelity"}
    report["visual_review_required"] = True
    (output / "export-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Exported {len(slides)} pages to {output}; visual review required")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("html", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--editable", action="store_true")
    parser.add_argument("--allow-hold", action="store_true")
    parser.add_argument("--browser-path", type=Path, default=Path(os.environ["YOUYOU_BROWSER_PATH"]) if os.environ.get("YOUYOU_BROWSER_PATH") else None)
    parser.add_argument("--from-frames", action="store_true", help="Use audited 1920x1080 PNGs already in output/frames; Python Playwright not required")
    args = parser.parse_args()
    export(args.html, args.output, args.editable, args.allow_hold, args.browser_path, args.from_frames)
